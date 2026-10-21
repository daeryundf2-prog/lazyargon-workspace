#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
crime_schedule_xlsx.py — 법원 제출용 범죄일람표·청취검토용 엑셀 자동 빌더 및 정합성 감사기

기능:
  1. 동아시아 글자 폭(W/F=2, 그 외=1) 기반 시각적 줄 수 계산 및 행 높이 자동 설정(글자 잘림 방지)
  2. 법원 제출용 표준 서식(맑은 고딕, 1행 제목 병합, 2행 #1F4E78 남색 헤더, A3 틀고정) 일괄 적용
  3. 유튜브 타임스탬프 바로가기 URL(`https://youtu.be/<id>?t=<sec>`) 생성
  4. 발언 원문(G열/transcript) ↔ 고소이유·쟁점(H열/issue) 키워드 정합성 감사
"""
from __future__ import annotations

import argparse
import json
import math
import re
import sys
import unicodedata
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

DEFAULT_HEADERS = [
    "연번", "방송일자", "영상시간\n(시작~종료)", "원본 영상 파일명",
    "원본 SHA-256\n(앞 16자리)", "클립 파일명",
    "피고소인 발언 원문 (서버 STT 전사)", "모욕·허위사실 요지 및 고소 이유",
    "유튜브 타임스탬프\n바로가기", "검증 상태",
]
DEFAULT_WIDTHS = {
    "A": 7.0, "B": 13.0, "C": 14.0, "D": 28.0, "E": 15.0,
    "F": 15.0, "G": 52.0, "H": 40.0, "I": 24.0, "J": 18.0,
}
_YT_ID_RE = re.compile(r"(?:v=|youtu\.be/|\[)([A-Za-z0-9_-]{11})(?:\]|&|\?|$)")


def calc_visual_lines(text: str, col_width: float = 20.0) -> int:
    """unicodedata.east_asian_width('W','F'=2, 나머지=1) 기반 시각적 줄 수를 계산한다."""
    if not text:
        return 1
    eff_w = max(4.0, (float(col_width) - 1.5) * 0.90)
    total = 0
    for para in str(text).splitlines() or [""]:
        w = sum(2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1 for ch in para)
        total += max(1, int(math.ceil(w / eff_w)))
    return max(1, total)


def recalc_sheet_row_heights(
    ws, start_row: int = 3, col_widths: dict | None = None, num_cols: int | None = None,
    min_pt: float = 30.0, max_pt: float = 409.0, line_pt: float = 15.2, pad_pt: float = 10.0,
) -> None:
    """openpyxl 워크시트의 열 너비 및 각 행 최대 시각적 줄 수를 계산해 행 높이를 설정한다."""
    from openpyxl.utils import get_column_letter
    ncols = num_cols or ws.max_column or 1
    for k, w in (col_widths or {}).items():
        ws.column_dimensions[get_column_letter(k) if isinstance(k, int) else str(k)].width = float(w)
    for r in range(start_row, (ws.max_row or start_row) + 1):
        max_lines = 1
        for c in range(1, ncols + 1):
            letter = get_column_letter(c)
            w = (col_widths or {}).get(letter) or (col_widths or {}).get(c) or ws.column_dimensions[letter].width or 18.0
            val = ws.cell(row=r, column=c).value
            if val is not None and not str(val).startswith("="):
                max_lines = max(max_lines, calc_visual_lines(str(val), float(w)))
        ws.row_dimensions[r].height = min(max_pt, max(min_pt, round(max_lines * line_pt + pad_pt, 1)))


def youtube_jump_url(video_id_or_url: str, start_sec: float, pre_roll: int = 2) -> str:
    """11자리 유튜브 ID를 추출해 https://youtu.be/<id>?t=<max(0, int(start_sec) - pre_roll)> 반환."""
    s = str(video_id_or_url).strip()
    m = _YT_ID_RE.search(s)
    vid = m.group(1) if m else (s if re.fullmatch(r"[A-Za-z0-9_-]{11}", s) else "")
    return f"https://youtu.be/{vid}?t={max(0, int(float(start_sec)) - int(pre_roll))}" if vid else ""


def style_court_schedule_sheet(
    ws, title_text: str, header_hex: str = "1F4E78", num_cols: int = 10, custom_widths: dict | None = None,
) -> None:
    """1행 제목 병합, 2행 남색 헤더, 본문 9pt 테두리·줄바꿈, A3 틀고정 및 행높이 재계산."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    widths = {**DEFAULT_WIDTHS, **(custom_widths or {})}
    thin = Side(style="thin", color="D0D7DE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.merge_cells(f"A1:{get_column_letter(num_cols)}1")
    ws["A1"] = title_text
    ws["A1"].font = Font(name="맑은 고딕", size=12, bold=True, color=header_hex)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 26
    hdr_fill, hdr_font = PatternFill("solid", fgColor=header_hex), Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
    for c in range(1, num_cols + 1):
        cell = ws.cell(row=2, column=c)
        cell.fill, cell.font, cell.border = hdr_fill, hdr_font, border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    body_font = Font(name="맑은 고딕", size=9)
    link_font = Font(name="맑은 고딕", size=9, color="0563C1", underline="single")
    for r in range(3, (ws.max_row or 2) + 1):
        for c in range(1, num_cols + 1):
            cell = ws.cell(row=r, column=c)
            is_link = str(cell.value or "").startswith(("http://", "https://", "=HYPERLINK"))
            cell.font, cell.border = (link_font if is_link else body_font), border
            cell.alignment = Alignment(horizontal="left" if c in (4, 7, 8) else "center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A3"
    recalc_sheet_row_heights(ws, start_row=3, col_widths=widths, num_cols=num_cols)


def audit_schedule_consistency(rows: list[dict], rules: dict[str, list[str]]) -> list[dict]:
    """발언 원문(transcript/G열)과 쟁점(issue/H열) 간 키워드 정합성 불일치 행을 반환한다."""
    mismatches: list[dict] = []
    for idx, row in enumerate(rows, 1):
        transcript = str(row.get("transcript") or row.get("G") or "")
        issue = str(row.get("issue") or row.get("H") or "")
        for issue_kw, req_words in rules.items():
            if issue_kw and issue_kw in issue and not any(w in transcript for w in req_words):
                mismatches.append({
                    "row": row.get("row", idx),
                    "issue_keyword": issue_kw,
                    "expected_any": list(req_words),
                    "transcript": transcript,
                    "issue": issue,
                })
    return mismatches


def _parse_rules(rule_args: list[str] | None) -> dict[str, list[str]]:
    rules: dict[str, list[str]] = {}
    for item in rule_args or []:
        if "=" in item:
            k, v = item.split("=", 1)
            rules[k.strip()] = [w.strip() for w in v.split(",") if w.strip()]
    return rules


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="법원 제출용 범죄일람표 엑셀 빌더 및 정합성 감사기")
    ap.add_argument("--recalc", help="기존 엑셀 파일(.xlsx) 행 높이 자동 재계산")
    ap.add_argument("--audit", help="엑셀(.xlsx) 또는 JSON 파일의 G/H열 키워드 정합성 감사")
    ap.add_argument("--rule", action="append", default=[], help="감사 규칙 '쟁점어=근거어1,근거어2' (복수 지정 가능)")
    ap.add_argument("--from-json", help="범죄일람표 행 데이터 JSON 경로")
    ap.add_argument("-o", "--output", help="생성할 엑셀(.xlsx) 출력 경로")
    ap.add_argument("--title", default="범죄일람표 (서버 STT 전사 및 원본 해시 대조필)", help="시트 1행 제목")
    args = ap.parse_args(argv)

    if args.recalc:
        import openpyxl

        p = Path(args.recalc)
        if not p.is_file():
            return 2
        wb = openpyxl.load_workbook(p)
        for ws in wb.worksheets:
            recalc_sheet_row_heights(ws, start_row=3)
        wb.save(args.output or p)
        return 0

    if args.audit:
        p = Path(args.audit)
        if not p.is_file():
            return 2
        rules = _parse_rules(args.rule)
        rows: list[dict] = []
        if p.suffix.lower() == ".json":
            raw = json.loads(p.read_text(encoding="utf-8"))
            rows = raw.get("rows", raw) if isinstance(raw, dict) else raw
        else:
            import openpyxl

            wb = openpyxl.load_workbook(p, data_only=True)
            for ws in wb.worksheets:
                for r in range(3, (ws.max_row or 2) + 1):
                    rows.append({"row": r, "G": ws.cell(r, 7).value or "", "H": ws.cell(r, 8).value or ""})
        bad = audit_schedule_consistency(rows, rules)
        print(json.dumps({"file": str(p), "mismatches": bad, "count": len(bad)}, ensure_ascii=False, indent=2))
        return 1 if bad else 0

    if args.from_json:
        import openpyxl

        src = Path(args.from_json)
        if not src.is_file() or not args.output:
            return 2
        data = json.loads(src.read_text(encoding="utf-8"))
        title = data.get("title", args.title) if isinstance(data, dict) else args.title
        headers = data.get("headers", DEFAULT_HEADERS) if isinstance(data, dict) else DEFAULT_HEADERS
        items = data.get("rows", data) if isinstance(data, dict) else data
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "범죄일람표"
        for c, h in enumerate(headers, 1):
            ws.cell(row=2, column=c, value=h)
        keys = ["no", "date", "time_range", "source_file", "sha256_prefix", "clip_file", "transcript", "issue", "jump_url", "status"]
        for r_idx, item in enumerate(items, 3):
            vals = [item.get(k, "") for k in keys] if isinstance(item, dict) else list(item)
            for c_idx, val in enumerate(vals[: len(headers)], 1):
                ws.cell(row=r_idx, column=c_idx, value=val)
        style_court_schedule_sheet(ws, title_text=title, num_cols=len(headers))
        wb.save(args.output)
        return 0

    ap.print_help()
    return 2


if __name__ == "__main__":
    sys.exit(main())
